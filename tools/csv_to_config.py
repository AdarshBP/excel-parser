"""Generate a configuration workbook from a CSV file's headers.

Reads the CSV, infers column types from the first N data rows, and writes
a config .xlsx with sheet_config, column_config, and target_config sheets.

    python csv_to_config.py data.csv                    # writes data_config.xlsx
    python csv_to_config.py data.csv -o my_config.xlsx  # custom output name

The generated config targets PostgreSQL by default (database = excel_parser).
"""
import argparse
import csv
import re
from pathlib import Path

import openpyxl


def _slugify(name: str) -> str:
    """Turn a header like 'Food Value (A)' into 'food_value_a'."""
    s = name.strip().lower()
    s = re.sub(r"[^a-z0-9]+", "_", s)
    s = s.strip("_")
    return s or "col"


def _infer_type(values: list[str]) -> str:
    """Guess the data_type from a sample of non-empty string values."""
    if not values:
        return "text"

    numeric, integer, date, datetime_ = 0, 0, 0, 0
    for v in values:
        v = v.strip()
        # Try numeric
        cleaned = v.replace(",", "").replace("₹", "").replace("$", "").strip()
        try:
            f = float(cleaned)
            numeric += 1
            if f == int(f) and "." not in cleaned:
                integer += 1
            continue
        except ValueError:
            pass
        # Try datetime (YYYY-MM-DD HH:MM:SS)
        if re.match(r"\d{4}-\d{2}-\d{2}\s+\d{2}:\d{2}(:\d{2})?", v):
            datetime_ += 1
            continue
        # Try date patterns
        if (re.match(r"\d{4}-\d{2}-\d{2}$", v) or
            re.match(r"\d{2}/\d{2}/\d{4}$", v) or
            re.match(r"(Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday),\s+\d{2}\s+\w+\s+\d{4}", v)):
            date += 1
            continue

    total = len(values)
    threshold = 0.8

    if (numeric / total) >= threshold:
        return "integer" if integer == numeric else "numeric"
    if (datetime_ / total) >= threshold:
        return "timestamp"
    if (date / total) >= threshold:
        return "date"
    return "text"


def _col_letter(index: int) -> str:
    """0-based index to Excel column letter (0=A, 25=Z, 26=AA)."""
    result = ""
    i = index
    while True:
        result = chr(65 + (i % 26)) + result
        i = i // 26 - 1
        if i < 0:
            break
    return result


def generate(csv_path: Path, output: Path = None,
             table_name: str = None, sample_rows: int = 100,
             database: str = "excel_parser", db_schema: str = None) -> Path:
    """Generate a config workbook from a CSV file."""
    with open(csv_path, newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        headers = next(reader)
        samples: list[list[str]] = [[] for _ in headers]
        for i, row in enumerate(reader):
            if i >= sample_rows:
                break
            for j, val in enumerate(row):
                if j < len(headers) and val.strip():
                    samples[j].append(val.strip())

    if not table_name:
        table_name = _slugify(csv_path.stem.split("_")[0] if "_" in csv_path.stem
                              else csv_path.stem)
    sheet_name = csv_path.stem

    wb = openpyxl.Workbook()

    # sheet_config
    sc = wb.active
    sc.title = "sheet_config"
    sc.append(["table_name", "sheet_name", "layout", "header_row",
               "data_start_row", "data_end_row", "active", "description", "domain", "notes"])
    sc.append([table_name, sheet_name, "table", 1, 2, None, "Y",
               f"Data from {csv_path.stem}", None,
               f"Auto-generated from {csv_path.name}"])

    # column_config
    cc = wb.create_sheet("column_config")
    cc.append(["table_name", "column_name", "source_ref", "data_type",
               "nullable", "column_order", "source_header", "script", "description", "unit"])
    for j, header in enumerate(headers):
        col_name = _slugify(header)
        col_letter = _col_letter(j)
        dtype = _infer_type(samples[j])
        # Auto-assign trim for text columns
        script = "trim" if dtype == "text" else None
        cc.append([table_name, col_name, f"col:{col_letter}", dtype,
                   "Y", j + 1, header, script, None, None])

    # target_config
    tc = wb.create_sheet("target_config")
    tc.append(["setting", "value"])
    tc.append(["target", "postgres"])
    tc.append(["database", database])
    if db_schema:
        tc.append(["db_schema", db_schema])
    tc.append(["id_type", "integer"])

    if output is None:
        output = csv_path.parent / f"{table_name}_config.xlsx"
    wb.save(output)
    return output


def main():
    ap = argparse.ArgumentParser(description="Generate config from CSV headers")
    ap.add_argument("csv", type=Path, help="source CSV file")
    ap.add_argument("-o", "--output", type=Path, default=None, help="output config path")
    ap.add_argument("--table", default=None, help="override table name")
    ap.add_argument("--database", default="excel_parser")
    ap.add_argument("--schema", default=None, help="PostgreSQL schema")
    ap.add_argument("--sample", type=int, default=100, help="rows to sample for type inference")
    args = ap.parse_args()

    if not args.csv.is_file():
        raise SystemExit(f"{args.csv} does not exist")

    out = generate(args.csv, args.output, args.table, args.sample,
                   args.database, args.schema)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
