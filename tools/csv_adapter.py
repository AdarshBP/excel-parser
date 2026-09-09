"""Read a CSV file through the same interface the parser uses for .xlsx.

A CsvWorkbook quacks like an openpyxl Workbook: it has .sheetnames and
supports wb[name] to get a worksheet. A CsvSheet quacks like an openpyxl
Worksheet: it has .max_row, .max_column, and .cell(row, col).value.

Every cell value is a string (or None for empty cells). The existing
values.py casting handles the conversion to numeric/date/etc, so nothing
is lost — no intermediate .xlsx, no floating-point surprises.

Usage:
    from csv_adapter import open_source
    wb = open_source(Path("data.csv"))   # CsvWorkbook
    wb = open_source(Path("data.xlsx"))  # openpyxl Workbook (unchanged)
"""
import csv
from pathlib import Path


class _Cell:
    """Minimal cell with a .value attribute."""
    __slots__ = ("value",)

    def __init__(self, value):
        self.value = value


class CsvSheet:
    """A CSV file presented as a single worksheet.

    Row and column numbers are 1-based, matching openpyxl convention.
    """

    def __init__(self, path: Path, sheet_name: str):
        self._rows: list[list[str]] = []
        self._max_column = 0
        self.title = sheet_name

        with open(path, newline="", encoding="utf-8-sig") as f:
            for row in csv.reader(f):
                self._rows.append(row)
                if len(row) > self._max_column:
                    self._max_column = len(row)

    @property
    def max_row(self) -> int:
        return len(self._rows)

    @property
    def max_column(self) -> int:
        return self._max_column

    def cell(self, row: int, column: int) -> _Cell:
        """Return a cell-like object. Row and column are 1-based."""
        if row < 1 or row > len(self._rows):
            return _Cell(None)
        r = self._rows[row - 1]
        if column < 1 or column > len(r):
            return _Cell(None)
        val = r[column - 1]
        return _Cell(val if val != "" else None)

    def iter_rows(self, values_only=False):
        """Yield rows like openpyxl does (used by validator.read_config)."""
        for r in self._rows:
            padded = r + [""] * (self._max_column - len(r))
            if values_only:
                yield tuple(v if v != "" else None for v in padded)
            else:
                yield tuple(_Cell(v if v != "" else None) for v in padded)


class CsvWorkbook:
    """A CSV file presented as a single-sheet workbook.

    The sheet name defaults to the filename without extension.
    """

    def __init__(self, path: Path):
        self._path = path
        self._name = path.stem
        self._sheet: CsvSheet | None = None

    def _ensure(self) -> CsvSheet:
        if self._sheet is None:
            self._sheet = CsvSheet(self._path, self._name)
        return self._sheet

    @property
    def sheetnames(self) -> list[str]:
        return [self._name]

    def __getitem__(self, name: str) -> CsvSheet:
        if name != self._name:
            raise KeyError(f"CSV workbook has no sheet '{name}' "
                           f"(only '{self._name}')")
        return self._ensure()

    def __contains__(self, name: str) -> bool:
        return name == self._name

    def close(self):
        pass


def is_csv(path: Path) -> bool:
    """Check if a path points to a CSV file."""
    return path.suffix.lower() == ".csv"


def open_source(path: Path):
    """Open a source file — returns CsvWorkbook for .csv, openpyxl for .xlsx."""
    if is_csv(path):
        return CsvWorkbook(path)
    import openpyxl
    return openpyxl.load_workbook(path, read_only=True, data_only=True)
